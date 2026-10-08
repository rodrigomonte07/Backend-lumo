#!/usr/bin/env python3
"""
Grava assets/line_breaks.json: para cada caixa de texto do template, as linhas EXATAS que o Canva usou (lidas do PDF).
Em tempo de geração viram <a:br/>, para o texto quebrar igual ao Canva em qualquer renderizador.
Uso: python3 tools/build_breaks.py assets/modelo_canva_referencia.pdf "modelo_rede_lumo (1).pptx" assets/line_breaks.json
"""
import json, os, re, sys
import pymupdf as fitz
from pptx import Presentation
sys.path.insert(0, os.path.dirname(__file__))
from build_pills import iter_abs

def norm(s): return re.sub(r'\s+', '', s)

def main(pdf, pptx, out):
    ref, prs = fitz.open(pdf), Presentation(pptx)
    res = {}
    for pn in range(2, 10):
        lines = []
        for b in ref[pn - 1].get_text('dict')['blocks']:
            for l in b.get('lines', []):
                t = ''.join(s['text'] for s in l['spans']).strip()
                if t:
                    lines.append((fitz.Rect(l['bbox']), t))
        for sh, l, t, w, h in iter_abs(prs.slides[pn - 1].shapes):
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            x0, y0, x1, y1 = l / 12700, t / 12700, (l + w) / 12700, (t + h) / 12700
            inside = [ln for ln in lines if x0 - 4 <= ln[0].x0 <= x1 + 4 and y0 - 8 <= ln[0].y0 <= y1 + 8 and ln[0].y1 <= y1 + 14]
            inside.sort(key=lambda ln: (round(ln[0].y0 / 3), ln[0].x0))
            if len(inside) >= 2 and norm(''.join(i[1] for i in inside)) == norm(sh.text_frame.text.replace('\v', '')):
                res.setdefault(str(pn), {})[str(sh.shape_id)] = [i[1] for i in inside]
    json.dump(res, open(out, 'w'), ensure_ascii=False, indent=1)
    print(sum(len(v) for v in res.values()), 'blocos de várias linhas')

if __name__ == '__main__':
    main(*sys.argv[1:4])
