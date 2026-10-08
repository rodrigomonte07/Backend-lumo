#!/usr/bin/env python3
"""
Extrai do PDF do Canva os "marca-textos" (fundos arredondados atrás de texto) que o Canva NÃO exporta
para o .pptx (ex.: o fundo pêssego de "O que enxergamos no ...", os fundos azul/rosa dos valores do slide 5,
os selos cinza de "Fechamento", etc.) e grava assets/pills.json.

Em tempo de geração, app/mail_merge.py desenha esses fundos como retângulos arredondados atrás do texto,
medindo o texto novo com a fonte real — assim o fundo acompanha o tamanho do valor/nome da escola.

Uso: python3 tools/build_pills.py assets/modelo_canva_referencia.pdf "modelo_rede_lumo (1).pptx" assets/pills.json
"""
import json, sys
import pymupdf as fitz
from pptx import Presentation

PAGES = range(2, 10)

def iter_abs(shapes, tx=(0.0, 0.0, 1.0, 1.0)):
    ox, oy, sx, sy = tx
    for sh in shapes:
        if sh.left is None:
            continue
        l, t, w, h = ox + sh.left * sx, oy + sh.top * sy, sh.width * sx, sh.height * sy
        yield sh, l, t, w, h
        if sh.shape_type == 6:
            x = sh._element.grpSpPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm')
            co, ce = x.find('{http://schemas.openxmlformats.org/drawingml/2006/main}chOff'), x.find('{http://schemas.openxmlformats.org/drawingml/2006/main}chExt')
            csx, csy = w / int(ce.get('cx')), h / int(ce.get('cy'))
            yield from iter_abs(sh.shapes, (l - int(co.get('x')) * csx, t - int(co.get('y')) * csy, csx, csy))

def norm(s): return ''.join(s.split()).lower()

def main(pdf, pptx, out):
    doc, prs = fitz.open(pdf), Presentation(pptx)
    pills = []
    for pn in PAGES:
        pg = doc[pn - 1]
        lines = []
        for b in pg.get_text('dict')['blocks']:
            for l in b.get('lines', []):
                t = ''.join(s['text'] for s in l['spans'])
                if t.strip():
                    lines.append((fitz.Rect(l['bbox']), t, l['spans'][0]['size']))
        shapes = [(x[0],) + tuple(v / 12700 for v in x[1:]) for x in iter_abs(prs.slides[pn - 1].shapes) if x[0].has_text_frame and x[0].text_frame.text.strip()]
        for d in pg.get_drawings():
            f = d.get('fill')
            if f is None or d.get('fill_opacity', 1) == 0:
                continue
            op = d.get('fill_opacity', 1) or 1
            r = d['rect']
            if r.width < 25 or r.height < 10 or r.x0 < 70 or r.x1 > 1330 or r.width > 1200:
                continue
            if not any(it[0] == 'c' for it in d['items']):       # sem curvas = é um shape que o pptx já tem
                continue
            inside = [x for x in lines if x[0].intersects(r) and (x[0] & r).get_area() > 0.7 * x[0].get_area()]
            if not inside:
                continue
            inside.sort(key=lambda x: (round(x[0].y0), x[0].x0))
            n, H = len(inside), r.height
            padx = inside[0][0].x0 - r.x0
            pad_r = r.x1 - max(x[0].x1 for x in inside)
            # shape do pptx que contém a primeira linha
            fl = inside[0]
            cx, cy = (fl[0].x0 + fl[0].x1) / 2, (fl[0].y0 + fl[0].y1) / 2
            cand = [s for s in shapes if s[1] - 2 <= cx <= s[1] + s[3] + 2 and s[2] - 2 <= cy <= s[2] + s[4] + 2]
            cand = [s for s in cand if norm(fl[1])[:8] in norm(s[0].text_frame.text)]
            if not cand:
                print('  (sem shape para)', pn, fl[1][:30]); continue
            sh, sl, st, sw, shh = min(cand, key=lambda s: s[3] * s[4])
            ppr = sh.text_frame.paragraphs[0]._p.find('{http://schemas.openxmlformats.org/drawingml/2006/main}pPr')
            algn = ppr.get('algn') if ppr is not None else 'l'
            pills.append({
                'slide': pn, 'shape_id': sh.shape_id,
                'color': '%02X%02X%02X' % tuple(round((c * op + (1 - op)) * 255) for c in f),
                'algn': algn or 'l', 'size_pt': round(fl[2], 1),
                'shape_left': round(sl, 2), 'shape_top': round(st, 2),
                'padx': round(padx, 2), 'pad_r': round(pad_r, 2),
                'lines': [{'text': x[1].strip(), 'x0': round(x[0].x0, 2), 'x1': round(x[0].x1, 2),
                           'top': round(r.y0 + i * H / n, 2), 'bottom': round(r.y0 + (i + 1) * H / n, 2)}
                          for i, x in enumerate(inside)],
            })
            print(pn, sh.shape_id, pills[-1]['color'], n, 'linha(s):', fl[1][:35])
    json.dump(pills, open(out, 'w'), ensure_ascii=False, indent=1)
    print(len(pills), 'fundos gravados em', out)

if __name__ == '__main__':
    main(*sys.argv[1:4])
