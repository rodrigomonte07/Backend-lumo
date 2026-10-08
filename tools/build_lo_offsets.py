#!/usr/bin/env python3
"""
Mede, em cada caixa de texto, o deslocamento vertical entre o LibreOffice e o PDF do Canva (comparando a LINHA DE BASE
de cada linha de texto) e grava assets/lo_offsets.json -> {slide: {shape_id: dy_pt}}. Duas passadas de calibração.
Usado só na conversão para PDF (app.mail_merge.tune_for_libreoffice): o PowerPoint posiciona esse texto como o Canva;
o LibreOffice distribui o espaço extra da entrelinha de outro jeito.

Uso (precisa de soffice):  python3 tools/build_lo_offsets.py
"""
import json, statistics, subprocess, sys, tempfile
from pathlib import Path
import pymupdf as fitz
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app import mail_merge as mm
from pptx import Presentation


def baselines(pg):
    out = []
    for b in pg.get_text('dict')['blocks']:
        for l in b.get('lines', []):
            t = ''.join(s['text'] for s in l['spans']).strip()
            if t:
                out.append((t, l['spans'][0]['origin'], l['spans'][0]['size']))
    return out


def main():
    ref = fitz.open(ROOT / 'assets/modelo_canva_referencia.pdf')
    tpl = ROOT / 'modelo_rede_lumo (1).pptx'
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / 'cr.pptx'
        fm = mm.load_json(ROOT / 'field_map.json')
        # dados = os próprios textos do template (Cristo Rei): mede o LibreOffice na MESMA configuração de produção
        data = {}
        for r in fm:
            if not r.get('fmt') and r['token'] not in data:
                data[r['token']] = r['find']
        data.update({'NOME_ESCOLA': 'Cristo Rei', 'PERCENTUAL_CASH_OUT': '47%', 'PERCENTUAL_CASH_IN': '53%'})
        mm.generate(tpl, fm, data, out, None, mm.new_report())
        total = {}
        for it in range(2):
            src = out
            if total:
                tmp_json = Path(tmp) / 'off.json'
                tmp_json.write_text(json.dumps(total))
                src = Path(tmp) / f'tuned{it}.pptx'
                mm.tune_for_libreoffice(out, src, tmp_json)
            subprocess.run(['soffice', f'-env:UserInstallation=file://{tmp}/lo{it}', '--headless', '--convert-to', 'pdf',
                            '--outdir', tmp, str(src)], check=True, capture_output=True, timeout=300)
            me = fitz.open(Path(tmp) / (src.stem + '.pdf'))
            prs = Presentation(str(tpl))
            for pn in range(1, 9):
                R, M = baselines(ref[pn]), baselines(me[pn])
                for sh, l, t, w, h in mm._iter_shapes_abs(prs.slides[pn].shapes):
                    if not (sh.has_text_frame and sh.text_frame.text.strip()):
                        continue
                    x0, y0, x1, y1 = l / 12700, t / 12700, (l + w) / 12700, (t + h) / 12700
                    shape_txt = ''.join(sh.text_frame.text.split())
                    ds = []
                    for txt, o, sz in R:
                        if not (x0 - 6 <= o[0] <= x1 + 6 and y0 - 12 <= o[1] <= y1 + 10) or ''.join(txt.split()) not in shape_txt:
                            continue
                        c = [m for m in M if m[0] == txt and abs(m[1][0] - o[0]) < 25 and abs(m[1][1] - o[1]) < 20]
                        if c:
                            ds.append(c[0][1][1] - o[1])
                    if ds:
                        dy = statistics.median(ds)
                        if abs(dy) >= 0.4:
                            d = total.setdefault(str(pn + 1), {})
                            d[str(sh.shape_id)] = round(d.get(str(sh.shape_id), 0) + dy, 2)
        table = total
    json.dump(table, open(ROOT / 'assets/lo_offsets.json', 'w'), indent=1)
    print(sum(len(v) for v in table.values()), 'caixas com correção')


if __name__ == '__main__':
    main()
