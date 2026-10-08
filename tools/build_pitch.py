#!/usr/bin/env python3
"""
O .pptx do Canva grava o espaçamento entre linhas (lnSpc/spcPts) 1–4% maior que o usado no design; em textos de
várias linhas o erro acumula. Mede o espaçamento REAL no PDF do Canva e grava assets/line_pitch.json.
Uso: python3 tools/build_pitch.py assets/modelo_canva_referencia.pdf "modelo_rede_lumo (1).pptx" assets/line_pitch.json
"""
import json, sys
import pymupdf as fitz
from pptx import Presentation
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

def main(pdf, pptx, out):
    ref, prs = fitz.open(pdf), Presentation(pptx)
    res = {}
    for pn in range(1, 9):
        lines = []
        for b in ref[pn].get_text('dict')['blocks']:
            for l in b.get('lines', []):
                t = ''.join(s['text'] for s in l['spans']).strip()
                if t:
                    lines.append((fitz.Rect(l['bbox']), l['spans'][0]['size']))
        for sh in prs.slides[pn].shapes:
            if not sh.has_text_frame:
                continue
            p0 = sh.text_frame.paragraphs[0]
            spc = p0._p.find('.//' + A + 'spcPts')
            if spc is None or not p0.runs or p0.runs[0].font.size is None:
                continue
            L, size = int(spc.get('val')) / 100, p0.runs[0].font.size.pt
            x0, y0 = sh.left / 12700, sh.top / 12700
            x1, y1 = x0 + sh.width / 12700, y0 + sh.height / 12700 + 30
            ls = sorted([l for l in lines if x0 - 5 <= l[0].x0 <= x1 and y0 - 8 <= l[0].y0 <= y1 and abs(l[1] - size) < 1],
                        key=lambda l: l[0].y0)
            if len(ls) < 2:
                continue
            pitch = (ls[-1][0].y0 - ls[0][0].y0) / (len(ls) - 1)
            if 0.94 <= pitch / L <= 1.02:
                res.setdefault(str(pn + 1), {})[str(sh.shape_id)] = round(pitch, 2)
    json.dump(res, open(out, 'w'), indent=1)
    print(sum(len(v) for v in res.values()), 'blocos com espaçamento medido')

if __name__ == '__main__':
    main(*sys.argv[1:4])
