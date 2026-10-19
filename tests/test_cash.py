"""Regra única de Cash-Out/Cash-In (com e sem assunção de dívida) — caso CEME Escola e Creche."""
from pathlib import Path
from pptx import Presentation
from app import mail_merge as mm

ROOT = Path(__file__).resolve().parent.parent
CEME = dict(
    VALOR_OPERACAO_ESCOLAR="R$ 455.740,34", VALOR_PATRIMONIO_IMOBILIARIO="R$ 677.600,00",
    VALOR_TOTAL_TRANSACAO="R$ 1.133.340,34", VALOR_ASSUNCAO_DIVIDA="R$ 222.362,19", ASSUNCAO_DIVIDA_ATIVA="sim",
    PERCENTUAL_CASH_OUT="55%", PERCENTUAL_CASH_IN="45%", PERCENTUAL_ENTRADA="10%",
    ESCOLA_QTD_PARCELAS_ENTRADA="1", IMOVEL_QTD_PARCELAS_ENTRADA="1",
    ESCOLA_QTD_PARCELAS_SALDO="48", IMOVEL_QTD_PARCELAS_SALDO="48",
    # valores "antigos" (dívida repartida proporcionalmente) que o CRM antigo mandava: devem ser ignorados
    ESCOLA_TOTAL_DESTINADO_VENDEDOR="R$ 201.478,07", IMOVEL_TOTAL_DESTINADO_VENDEDOR="R$ 299.559,92",
)


def test_ceme_com_assuncao_de_divida():
    d = mm.derive_fields(CEME, {})
    assert d["VALOR_SUBTOTAL_TRANSACAO"] == "R$ 1.133.340,34"
    assert (d["S5_SUBTOTAL"], d["S5_DIVIDA"], d["S5_EXPOSICAO"]) == ("R$ 1.133.340,34", "R$ 222.362,19", "R$ 910.978,15")
    assert d["VALOR_CASH_OUT"] == "R$ 501.037,98" and d["VALOR_CASH_IN"] == "R$ 409.940,17"
    assert d["ESCOLA_TOTAL_DESTINADO_VENDEDOR"] == "R$ 128.357,98" and d["ESCOLA_VALOR_CASH_IN"] == "R$ 105.020,17"
    assert d["ESCOLA_VALOR_ENTRADA_TOTAL"] == "R$ 12.835,80" and d["ESCOLA_VALOR_PARCELA_SALDO"] == "R$ 2.406,71"
    assert d["IMOVEL_TOTAL_DESTINADO_VENDEDOR"] == "R$ 372.680,00" and d["IMOVEL_VALOR_CASH_IN"] == "R$ 304.920,00"
    assert d["IMOVEL_VALOR_ENTRADA_TOTAL"] == "R$ 37.268,00" and d["IMOVEL_VALOR_PARCELA_SALDO"] == "R$ 6.987,75"
    # conferências
    c = lambda k: mm._dec(d[k])
    assert c("ESCOLA_TOTAL_DESTINADO_VENDEDOR") + c("IMOVEL_TOTAL_DESTINADO_VENDEDOR") == c("VALOR_CASH_OUT")
    assert c("VALOR_CASH_OUT") + c("VALOR_CASH_IN") + mm._dec("R$ 222.362,19") == c("VALOR_SUBTOTAL_TRANSACAO")


def test_sem_assuncao_nada_e_descontado():
    dados = {k: v for k, v in CEME.items() if k not in ("VALOR_ASSUNCAO_DIVIDA", "ASSUNCAO_DIVIDA_ATIVA")}
    d = mm.derive_fields(dados, {})
    assert not d["S5_DIVIDA_ATIVA"] and "S5_DIVIDA" not in d
    assert d["S5_BASE_PERCENTUAL"] == "R$ 1.133.340,34"
    assert d["VALOR_CASH_OUT"] == "R$ 623.337,19"          # 1.133.340,34 x 55%
    assert d["ESCOLA_TOTAL_DESTINADO_VENDEDOR"] == "R$ 250.657,19"   # 455.740,34 x 55%
    assert d["IMOVEL_TOTAL_DESTINADO_VENDEDOR"] == "R$ 372.680,00"


def test_ultima_parcela_absorve_centavos():
    p, ultima = mm._split_parcelas(mm._dec("115522.18"), 48)
    assert (p, ultima) == (mm._dec("2406.71"), mm._dec("2406.81"))
    assert p * 47 + ultima == mm._dec("115522.18")


def test_slides_5_e_6_ceme(tmp_path):
    out = tmp_path / "o.pptx"
    mm.generate(ROOT / "modelo_rede_lumo (1).pptx", mm.load_json(ROOT / "field_map.json"), CEME, out)
    prs = Presentation(str(out))
    texts = {n: " | ".join(sh.text_frame.text for sh, *_ in mm._iter_shapes_abs(prs.slides[n - 1].shapes) if sh.has_text_frame)
             for n in (5, 6)}
    t5, t6 = texts[5], texts[6]
    for v in ("R$ 910.978,15", "R$ 501.037,98", "R$ 409.940,17", "R$ 222.362,19", "R$ 1.133.340,34", "55%", "45%"):
        assert v in t5, v
    for v in ("R$ 128.357,98", "R$ 372.680,00", "R$ 12.835,80", "R$ 37.268,00", "R$ 2.406,71", "R$ 6.987,75", "R$ 501.037,98"):
        assert v in t6, v
    for velho in ("201.478,07", "299.559,92"):
        assert velho not in t5 and velho not in t6
