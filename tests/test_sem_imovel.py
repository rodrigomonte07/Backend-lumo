"""Operação sem imóvel: slides 4, 6 e 7 sem os blocos do imóvel."""
from pathlib import Path
from pptx import Presentation
from app import mail_merge as mm

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "modelo_rede_lumo (1).pptx"
DADOS = dict(
    NOME_ESCOLA="Colégio Oliveira Mesquita", RAZAO_SOCIAL_ESCOLA="Colégio Oliveira Mesquita",
    PERCENTUAL_PARTICIPACAO_SOCIETARIA="100%", VALOR_OPERACAO_ESCOLAR="R$ 410.000,00", VALOR_TOTAL_TRANSACAO="R$ 410.000,00",
    PERCENTUAL_CASH_OUT="100%", VALOR_CASH_OUT="R$ 410.000,00", PCT_PARTICIPACAO_MANTENEDOR="0%", PCT_PARTICIPACAO_ADQUIRIDA="100%",
    ESCOLA_QTD_PARCELAS_ENTRADA="2", ESCOLA_QTD_PARCELAS_SALDO="36", PERCENTUAL_ENTRADA="12%",
    ASSUNCAO_DIVIDA_ATIVA="nao", VALOR_ASSUNCAO_DIVIDA="R$ 0,00",
)


def _texts(prs, n):
    return " | ".join(s.text_frame.text for s in prs.slides[n - 1].shapes if s.has_text_frame)


def _gerar(tmp_path, dados):
    out = tmp_path / "o.pptx"
    mm.generate(TEMPLATE, mm.load_json(ROOT / "field_map.json"), dados, out)
    return Presentation(str(out))


def test_deteccao():
    assert mm._is_sem_imovel(DADOS)
    assert mm._is_sem_imovel({**DADOS, "TIPO_OPERACAO": "sem_imovel", "VALOR_PATRIMONIO_IMOBILIARIO": "R$ 100,00"})
    assert not mm._is_sem_imovel({**DADOS, "TIPO_OPERACAO": "com_imovel"})
    assert not mm._is_sem_imovel({**DADOS, "VALOR_PATRIMONIO_IMOBILIARIO": "R$ 2.000.000,00"})


def test_slides_sem_blocos_do_imovel(tmp_path):
    prs = _gerar(tmp_path, DADOS)
    for n in (4, 6, 7):
        t = _texts(prs, n).lower()
        assert "imóvel" not in t and "imobili" not in t and "patrimônio" not in t and "aluguel" not in t, (n, t)
    assert "410.000,00" in _texts(prs, 4) and "49.200,00" in _texts(prs, 6) and "R$ 10.022,22" in _texts(prs, 7)
    assert "R$ 24.600,00" in _texts(prs, 7)


def test_com_imovel_nao_muda(tmp_path):
    t = _texts(_gerar(tmp_path, {**DADOS, "VALOR_PATRIMONIO_IMOBILIARIO": "R$ 500.000,00"}), 4)
    assert "PATRIMÔNIO IMOBILIÁRIO" in t
